# -*- coding: utf-8 -*-
"""Builds the Vacuum Control (QR Housekeeping) platform presentation (PPTX)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "mobile")  # real mobile-viewport captures (phone screen)

# --- Palette ---
BLUE   = RGBColor(0x15, 0x60, 0xA8)
DARK   = RGBColor(0x0F, 0x1F, 0x30)
INK    = RGBColor(0x1F, 0x2A, 0x37)
GRAY   = RGBColor(0x64, 0x74, 0x8B)
LIGHT  = RGBColor(0xEE, 0xF1, 0xF5)
CARD   = RGBColor(0xFF, 0xFF, 0xFF)
GREEN  = RGBColor(0x16, 0xA3, 0x4A)
AMBER  = RGBColor(0xF5, 0x9E, 0x0B)
RED    = RGBColor(0xEF, 0x44, 0x44)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def rect(s, x, y, w, h, color, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp

def rrect(s, x, y, w, h, color, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try: shp.adjustments[0] = 0.06
    except Exception: pass
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=1.0):
    """runs: list of paragraphs; each paragraph is list of (txt, size, color, bold)."""
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = space; p.space_after = Pt(4)
        for (txt, size, color, bold) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
            r.font.name = "Segoe UI"
    return tb

def image_fit(s, path, box_x, box_y, box_w, box_h, shadow=True):
    """Places the image (contain) inside the box, centered, with a frame."""
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih; box_ar = box_w / box_h
    if ar > box_ar:
        w = box_w; h = int(box_w / ar)
    else:
        h = box_h; w = int(box_h * ar)
    x = box_x + (box_w - w) // 2
    y = box_y + (box_h - h) // 2
    if shadow:
        # sombra rectangular (esquinas cuadradas) para que coincida con la imagen,
        # que también es un rectángulo de esquinas cuadradas.
        rect(s, x + Emu(20000), y + Emu(20000), w, h, RGBColor(0xD5,0xDB,0xE3))
    pic = s.shapes.add_picture(path, x, y, w, h)
    pic.line.color.rgb = RGBColor(0xE5,0xE9,0xF0); pic.line.width = Pt(1)
    return pic

def header(s, kicker, title):
    rect(s, 0, 0, SW, Inches(0.12), BLUE)
    text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.4),
         [[(kicker, 12, BLUE, True)]])
    text(s, Inches(0.6), Inches(0.65), Inches(12.1), Inches(0.7),
         [[(title, 27, DARK, True)]])

def sp(inches): return Inches(inches)

# ============ 1. COVER ============
s = slide(); bg(s, DARK)
rect(s, 0, 0, SW, SH, DARK)
rect(s, 0, Inches(6.9), SW, Inches(0.6), BLUE)
text(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(0.5),
     [[("VALENTISOFT  ·  SaaS PLATFORM", 14, RGBColor(0x8F,0xB4,0xDA), True)]])
text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.6),
     [[("Vacuum Control", 54, CARD, True)],
      [("QR-based housekeeping equipment control", 26, RGBColor(0xC7,0xD6,0xE8), False)]])
text(s, Inches(0.9), Inches(4.5), Inches(11.5), Inches(1.4),
     [[("Housekeepers report the status of every vacuum by scanning a QR code —", 16, RGBColor(0xB6,0xC6,0xDA), False)],
      [("no apps, no logins. Admins get automatic KPIs and reports.", 16, RGBColor(0xB6,0xC6,0xDA), False)]])
text(s, Inches(0.9), Inches(6.98), Inches(11.5), Inches(0.45),
     [[("Client: MasterCorp   ·   mastercorp.valentisoft.com", 13, CARD, True)]], anchor=MSO_ANCHOR.MIDDLE)

# ============ 2. THE CHALLENGE ============
s = slide(); bg(s, LIGHT); header(s, "CONTEXT", "The challenge: reporting and tracking equipment is scattered and manual")
cols = [
    ("Manual reporting", "The housekeeper reports verbally or on paper when a vacuum fails. It gets lost, forgotten, or reaches maintenance late."),
    ("No visibility", "The supervisor has no real-time view of which equipment is operational, faulty, or out of service at each site."),
    ("Reports by hand", "Consolidating the day's status and putting together the email for management is time-consuming and error-prone."),
]
x = Inches(0.6)
for (t, d) in cols:
    c = rrect(s, x, Inches(1.9), Inches(3.9), Inches(3.2), CARD)
    rect(s, x, Inches(1.9), Inches(3.9), Inches(0.14), RED)
    text(s, x+sp(0.25), Inches(2.25), Inches(3.4), Inches(0.6), [[(t, 18, INK, True)]])
    text(s, x+sp(0.25), Inches(2.95), Inches(3.4), Inches(2.0), [[(d, 14, GRAY, False)]], space=1.1)
    x += Inches(4.13)
text(s, Inches(0.6), Inches(5.6), Inches(12), Inches(1.2),
     [[("The result: ", 17, INK, True), ("equipment down longer, decisions made blind, and man-hours wasted on repetitive administrative tasks.", 17, INK, False)]], space=1.1)

# ============ 3. THE SOLUTION (how it works) ============
s = slide(); bg(s, LIGHT); header(s, "THE SOLUTION", "A simple flow: scan → report → visualize")
steps = [
    ("1", "QR on every unit", "Each vacuum carries a QR sticker with its code (VAC-001…). Generated and printed from the panel."),
    ("2", "The housekeeper scans", "Opens a web page (no app, no login), picks their name and the equipment status. 30 seconds."),
    ("3", "Logged instantly", "The report is saved with time, employee and site. It feeds the KPIs automatically."),
    ("4", "The admin decides", "Dashboard with status distribution, most-reported units and an automatic daily email to management."),
]
x = Inches(0.6)
for (n, t, d) in steps:
    c = rrect(s, x, Inches(2.0), Inches(2.95), Inches(3.4), CARD)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x+sp(0.25), Inches(2.25), Inches(0.7), Inches(0.7))
    circ.fill.solid(); circ.fill.fore_color.rgb = BLUE; circ.line.fill.background(); circ.shadow.inherit=False
    tf = circ.text_frame; tf.word_wrap=False; p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=n; r.font.size=Pt(22); r.font.bold=True; r.font.color.rgb=CARD
    text(s, x+sp(0.25), Inches(3.15), Inches(2.5), Inches(0.6), [[(t, 16, BLUE, True)]])
    text(s, x+sp(0.25), Inches(3.75), Inches(2.5), Inches(1.6), [[(d, 13, GRAY, False)]], space=1.1)
    if n != "4":
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x+sp(2.98), Inches(3.35), Inches(0.28), Inches(0.28))
        ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor(0xB6,0xC6,0xDA); ar.line.fill.background(); ar.shadow.inherit=False
    x += Inches(3.05)
text(s, Inches(0.6), Inches(5.75), Inches(12), Inches(0.8),
     [[("Multi-site and multi-client: ", 16, INK, True),
       ("each company gets its own subdomain (mastercorp.valentisoft.com) and each site its own equipment, QRs and KPIs.", 16, INK, False)]], space=1.1)

# ============ Feature slides (screenshot + how to use) ============
def feature(kicker, title, img, what, steps_list):
    s = slide(); bg(s, LIGHT); header(s, kicker, title)
    # text on the left
    text(s, Inches(0.6), Inches(1.75), Inches(4.5), Inches(0.9), [[(what, 15, INK, False)]], space=1.12)
    runs = []
    for st in steps_list:
        runs.append([("•  ", 14, BLUE, True), (st, 14, INK, False)])
    text(s, Inches(0.6), Inches(2.9), Inches(4.6), Inches(4.2), runs, space=1.25)
    # screenshot on the right
    image_fit(s, os.path.join(SHOTS, img), Inches(5.35), Inches(1.7), Inches(7.5), Inches(5.4))
    return s

feature("ADMIN", "Sites panel", "02-admin-sites.jpg",
        "The central view of all the client's sites, each with its equipment count and its recipients.",
        ["List of sites with their code, number of vacuums and report recipients.",
         "3 shortcuts per row: ⚙ configure · 📊 KPIs · ▦ QR codes.",
         "Each site has a unique slug used in the public URLs (/XjUS3/…).",
         "Responsive: adapts to mobile by hiding secondary columns."])

feature("ADMIN", "Configure a site and its equipment", "03-site-config.jpg",
        "This is where you set the site code, the recipients' emails (CC) and its vacuums.",
        ["Edit code and recipients (email validation with an example if wrong).",
         "Add vacuums: numbered automatically (VAC-001, 002…).",
         "Confirmation before creating equipment to avoid mistakes.",
         "Shortcuts to the site's QR codes and Dashboard."])

feature("ADMIN", "QR code generator", "04-qr-generator.jpg",
        "Generates each unit's QR with the client's logo in the center, ready to print and stick.",
        ["One QR per unit with the public URL ready to scan.",
         "Download one QR (PNG) or ALL of the site's in a printable PDF sheet.",
         "Fixed housekeeping-room QR to report 'equipment unavailable'.",
         "Copy-URL button and bilingual interface (EN/ES)."])

feature("ADMIN", "Per-site KPI dashboard", "05-kpi-dashboard.jpg",
        "Real-time indicators: status distribution, most-reported units and availability.",
        ["Status donut: Operational / With faults / Out of service / Unavailable.",
         "Ranking of most-reported units and availability %.",
         "Daily / Weekly / Monthly selector with a trend chart.",
         "Button to download the report as PDF."])

feature("ADMIN", "Report detail", "06-kpi-detail.jpg",
        "Every check-in with employee, unit, time/date, status and note. Full traceability.",
        ["Detail table with the record of every report in the period.",
         "Daily shows the time; weekly/monthly, the date.",
         "Separate 'equipment unavailable' section.",
         "Everything feeds the automatic daily email to recipients."])

feature("HOUSEKEEPER", "Scan a unit (no login)", "07-equipment-intro.jpg",
        "What the housekeeper sees when scanning a vacuum's QR. Simple, in their language, nothing to install.",
        ["Scan with the phone camera (iPhone or Android).",
         "See the client's branding and the unit (Vacuum · VAC-001).",
         "One button: 'Report'. EN/ES language switch.",
         "Works in any mobile browser."])

feature("HOUSEKEEPER", "Report equipment status", "08-report-form.jpg",
        "A 30-second form: who is reporting, what state the unit is in, and an optional note.",
        ["Employee autocomplete (pick your name from the list).",
         "3 color-coded states: Operational · With faults · Out of service.",
         "Optional note and photos when there is a problem.",
         "On submit, the record feeds the KPIs instantly."])

feature("HOUSEKEEPER", "Report 'equipment unavailable'", "09-housekeeping-intro.jpg",
        "A fixed QR in the housekeeping room to flag when NO vacuum is available.",
        ["Fixed QR posted in the room (not tied to a unit).",
         "The housekeeper scans and reports the shortage.",
         "Useful to detect equipment shortages per shift/floor.",
         "Counted separately in the dashboard."])

feature("HOUSEKEEPER", "Equipment-unavailable form", "10-housekeeping-form.jpg",
        "Employee, equipment type and comment. As simple as the status report.",
        ["Employee autocomplete.",
         "Equipment type (Vacuum) pre-selected.",
         "Optional comment (e.g. 'no vacuums on floor 3').",
         "Recorded with time and site."])

# ============ ROI: MAN-HOUR SAVINGS ============
s = slide(); bg(s, DARK); rect(s, 0, 0, SW, Inches(0.12), BLUE)
text(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.4), [[("THE VALUE", 12, RGBColor(0x8F,0xB4,0xDA), True)]])
text(s, Inches(0.6), Inches(0.7), Inches(12.1), Inches(0.7), [[("Man-hour savings and faster tasks", 27, CARD, True)]])
# 3 metric cards
metrics = [
    ("~90%", "less time per report", "From ~7 min (paper + manual entry)\nto ~30 sec scanning a QR."),
    ("~40 min", "per site per day", "The daily report is no longer built by hand:\nit is generated and sent on its own at 10 am."),
    ("~48 h", "saved per site per month", "Adding up reports + consolidation.\nWith 5 sites: ~240 h/month (~30 man-days)."),
]
x = Inches(0.6)
for (big, lbl, d) in metrics:
    c = rrect(s, x, Inches(1.75), Inches(3.9), Inches(3.0), RGBColor(0x14,0x2A,0x40))
    text(s, x+sp(0.3), Inches(2.05), Inches(3.4), Inches(0.9), [[(big, 40, RGBColor(0x5A,0xC8,0x8A), True)]])
    text(s, x+sp(0.3), Inches(2.95), Inches(3.4), Inches(0.5), [[(lbl, 15, CARD, True)]])
    text(s, x+sp(0.3), Inches(3.5), Inches(3.4), Inches(1.2), [[(d, 12.5, RGBColor(0xB6,0xC6,0xDA), False)]], space=1.1)
    x += Inches(4.13)
text(s, Inches(0.6), Inches(5.05), Inches(12.1), Inches(0.4),
     [[("Plus value that is hard to measure in hours:", 15, CARD, True)]])
extra = [
    "Faster maintenance: faulty equipment is detected instantly, not at the end of the shift.",
    "Data-driven decisions: real-time visibility of availability per site.",
    "Fewer errors: no lost paper and no manual transcription.",
    "Scales with no extra management cost: adding sites or equipment is immediate.",
]
runs = [[("•  ", 13, RGBColor(0x5A,0xC8,0x8A), True), (e, 13, RGBColor(0xC7,0xD6,0xE8), False)] for e in extra]
text(s, Inches(0.6), Inches(5.5), Inches(12.1), Inches(1.6), runs, space=1.2)
text(s, Inches(0.6), Inches(7.05), Inches(12), Inches(0.35),
     [[("Illustrative estimates based on a typical manual process; they vary with volume and operation.", 10, GRAY, False)]])

# ============ BENEFITS / CLOSING ============
s = slide(); bg(s, LIGHT); header(s, "SUMMARY", "Why Vacuum Control")
benefits = [
    ("Zero friction for the operator", "No apps, no logins: scan and report in seconds, in their language."),
    ("Real-time visibility", "Per-site KPIs, an automatic daily email and full traceability."),
    ("Multi-site and multi-client", "One subdomain per company, unlimited sites, each with its own equipment."),
    ("Secure and self-service", "Encrypted passwords, code-based recovery, admin roles."),
    ("Print-ready", "QRs with the client's branding, PDF sheets and a fixed housekeeping QR."),
    ("Measurable savings", "Fewer man-hours on reporting and consolidation; faster maintenance."),
]
x0, y0 = Inches(0.6), Inches(1.9)
for i, (t, d) in enumerate(benefits):
    col = i % 2; row = i // 2
    x = x0 + Inches(6.25)*col; y = y0 + Inches(1.55)*row
    c = rrect(s, x, y, Inches(5.9), Inches(1.35), CARD)
    rect(s, x, y, Inches(0.14), Inches(1.35), BLUE)
    text(s, x+sp(0.35), y+sp(0.15), Inches(5.3), Inches(0.5), [[(t, 16, BLUE, True)]])
    text(s, x+sp(0.35), y+sp(0.62), Inches(5.3), Inches(0.7), [[(d, 13, GRAY, False)]], space=1.05)

# ============ APPENDIX: ACCESS & SECURITY (at the end) ============
feature("ACCESS & SECURITY", "Secure per-client sign-in", "01-login.jpg",
        "Each client signs in through its subdomain with its own branding. Encrypted passwords (PBKDF2) and a persistent session.",
        ["The admin opens mastercorp.valentisoft.com and sees the login with their company logo.",
         "Enters email and password; long session (does not expire in normal use).",
         "Admin roles at the tenant level (e.g. Chris David, Ramces Rodriguez).",
         "Session invalidated when the password changes (account protection)."])

feature("ACCESS & SECURITY", "Password recovery", "11-forgot.jpg",
        "If an admin forgets their password, they recover it on their own, without relying on support.",
        ["'Forgot password?' → enter your email.",
         "Receive a 6-digit code by email (expires in 15 min).",
         "Enter the code and set the new password.",
         "Per-IP rate limiting against brute force; does not reveal whether the email exists."])

out = os.path.join(HERE, "VacuumControl-Presentation.pptx")
prs.save(out)
print("PPTX generated:", out, "· slides:", len(prs.slides._sldIdLst))
